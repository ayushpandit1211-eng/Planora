from contextlib import closing
import csv
import sqlite3
import uuid
import zipfile

import openpyxl
import pyarrow as pa
import pyarrow.parquet as pq
import pytest

from app.dataset_ingestion import DatasetIngestionService, analyze_file
from app.errors import QuerySentinelError
from app.structured_formats import convert_file, inspect_file
from tests.upgrade_helpers import admin_exec, requires_admin_db

CAP = 10 * 1024 * 1024


def fixture_file(tmp_path, kind):
    path = tmp_path / ("sample." + kind)
    if kind == "xlsx":
        book = openpyxl.Workbook()
        sheet = book.active
        sheet.title = "Sales & returns"
        sheet.append(["id", "name", "amount"])
        sheet.append([1, "Delhi", 12.5])
        sheet.append([2, "Pune", 20.0])
        book.save(path)
        book.close()
    elif kind == "parquet":
        pq.write_table(pa.table({"id": [1, 2], "name": ["Delhi", "Pune"], "amount": [12.5, 20.0]}), path)
    elif kind == "sqlite":
        with closing(sqlite3.connect(path)) as db:
            db.execute('CREATE TABLE "Sales & returns" (id integer, name text, amount real)')
            db.executemany('INSERT INTO "Sales & returns" VALUES(?,?,?)', [(1, "Delhi", 12.5), (2, "Pune", 20.)])
            db.commit()
    elif kind == "zip":
        with zipfile.ZipFile(path, "w") as archive:
            archive.writestr("data/sales.csv", "id,name,amount\n1,Delhi,12.5\n2,Pune,20.0\n")
    else:
        path.write_text("id\tname\tamount\n1\tDelhi\t12.5\n2\tPune\t20.0\n", encoding="utf-8")
    return path


@pytest.mark.parametrize("kind", ["xlsx", "parquet", "sqlite", "zip"])
def test_conversion_and_inference(tmp_path, kind):
    path = fixture_file(tmp_path, kind)
    before = path.read_bytes()
    output = tmp_path / "converted.csv"
    filename, notes = convert_file(path, path.name, output, CAP)
    a = analyze_file(output, filename)
    assert a.rows_read == 2
    assert a.column_types() == {"id": "integer", "name": "text", "amount": "numeric"}
    assert path.read_bytes() == before
    assert notes


def test_sheet_and_archive_selection(tmp_path):
    path = fixture_file(tmp_path, "xlsx")
    book = openpyxl.load_workbook(path)
    book.create_sheet("Other").append(["different"])
    book.save(path)
    book.close()
    assert inspect_file(path, path.name, CAP)["tables"] == ["Sales & returns", "Other"]
    with pytest.raises(QuerySentinelError):
        convert_file(path, path.name, tmp_path / "out.csv", CAP)
    archive = tmp_path / "multi.zip"
    with zipfile.ZipFile(archive, "w") as z:
        z.write(path, "nested/book.xlsx")
        z.writestr("a.tsv", "a\tb\n1\t2\n")
    assert len(inspect_file(archive, archive.name, CAP)["members"]) == 2
    assert len(inspect_file(archive, archive.name, CAP, "nested/book.xlsx")["tables"]) == 2
    convert_file(archive, archive.name, tmp_path / "out.csv", CAP, "Sales & returns", "nested/book.xlsx")
    assert analyze_file(tmp_path / "out.csv", "out.csv").rows_read == 2


@pytest.mark.parametrize("entry", ["../escape.csv", "/absolute.csv", "C:/outside.csv"])
def test_archive_paths_rejected(tmp_path, entry):
    path = tmp_path / "bad.zip"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr(entry, "a\n1\n")
    with pytest.raises(QuerySentinelError):
        inspect_file(path, path.name, CAP)


def test_expansion_and_formula_rejected(tmp_path):
    path = tmp_path / "bomb.zip"
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("big.csv", "x" * 10000)
    with pytest.raises(QuerySentinelError):
        inspect_file(path, path.name, 500)
    path = fixture_file(tmp_path, "xlsx")
    book = openpyxl.load_workbook(path)
    book.active["D2"] = "=1+1"
    book.save(path)
    book.close()
    with pytest.raises(QuerySentinelError, match="formulas"):
        convert_file(path, path.name, tmp_path / "out.csv", CAP)
    assert not (tmp_path / "out.csv").exists()


def test_sqlite_views_are_not_imported(tmp_path):
    path = fixture_file(tmp_path, "sqlite")
    with closing(sqlite3.connect(path)) as db:
        db.execute("CREATE VIEW unsafe AS SELECT load_extension('bad')")
        db.commit()
    assert inspect_file(path, path.name, CAP)["tables"] == ["Sales & returns"]
    with pytest.raises(QuerySentinelError):
        convert_file(path, path.name, tmp_path / "out.csv", CAP, "unsafe")


@requires_admin_db
@pytest.mark.asyncio
@pytest.mark.parametrize("kind", ["xlsx", "parquet", "sqlite", "zip", "tsv"])
async def test_real_database_import(tmp_path, kind):
    svc = DatasetIngestionService()
    dataset = "format_" + uuid.uuid4().hex[:10]
    path = fixture_file(tmp_path, kind)
    preview = await svc.preview(path, path.name, dataset_name=dataset, session_id="format-test-session")
    try:
        assert preview["rows_read"] == 2
        result = await svc.commit(preview["preview_id"], "format-test-session")
        assert result["rows_loaded"] == 2
        rows = admin_exec(f'SELECT name, amount FROM qs_dynamic."{dataset}" ORDER BY id')
        assert rows[0][0] == "Delhi" and float(rows[1][1]) == 20
    finally:
        # Existing service removes its test-owned dataset and metadata transactionally.
        if "result" in locals():
            await svc.delete_source(result["data_source_id"], "format-test-session")
        svc.discard(preview["preview_id"])
