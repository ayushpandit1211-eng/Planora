"""Bounded, read-only adapters to the existing CSV ingestion pipeline."""
from __future__ import annotations

import contextlib
import csv
import datetime as dt
import decimal
import json
import math
import sqlite3
import stat
import tempfile
import time
import zipfile
from pathlib import Path, PurePosixPath

from app.errors import ErrorCategory, QuerySentinelError

TEXT_EXT = {".csv", ".tsv", ".txt", ".json", ".jsonl", ".ndjson"}
STRUCTURED_EXT = {".xlsx", ".parquet", ".sqlite", ".sqlite3", ".db", ".zip"}
MAX_ROWS, MAX_COLUMNS = 1_000_000, 512


def reject(message):
    raise QuerySentinelError(ErrorCategory.UPLOAD_REJECTED, message, stage="upload")


def checked_zip(path, cap):
    archive = zipfile.ZipFile(path)
    try:
        entries = archive.infolist()
        if len(entries) > 2048 or sum(i.file_size for i in entries) > cap:
            reject("Archive expands beyond the upload limit or contains too many entries")
        for item in entries:
            name = PurePosixPath(item.filename.replace("\\", "/"))
            if name.is_absolute() or ".." in name.parts or ":" in item.filename:
                reject("Archive contains an unsafe path")
            if stat.S_ISLNK(item.external_attr >> 16) or item.flag_bits & 1:
                reject("Encrypted archives and symbolic links are not supported")
        return archive
    except BaseException:
        archive.close()
        raise


@contextlib.contextmanager
def source_file(path: Path, filename: str, cap: int, member: str | None = None):
    if path.stat().st_size > cap:
        reject("File exceeds upload limit")
    if Path(filename).suffix.lower() != ".zip":
        if member:
            reject("Archive member is only valid for ZIP files")
        yield path, filename
        return
    with checked_zip(path, cap) as archive:
        choices = [i.filename for i in archive.infolist() if not i.is_dir()
                   and Path(i.filename).suffix.lower() in (TEXT_EXT | STRUCTURED_EXT) - {".zip"}]
        if len(choices) != len(set(choices)):
            reject("Archive contains duplicate filenames")
        selected = member or (choices[0] if len(choices) == 1 else None)
        if selected not in choices:
            reject("Select one supported dataset file from the ZIP archive")
        with tempfile.TemporaryDirectory(prefix="planora_zip_", dir=path.parent) as directory:
            target = Path(directory) / ("member" + Path(selected).suffix.lower())
            with archive.open(selected) as src, target.open("wb") as dst:
                copied = 0
                while block := src.read(65536):
                    copied += len(block)
                    if copied > cap:
                        reject("Archive member expands beyond the upload limit")
                    dst.write(block)
            yield target, selected


@contextlib.contextmanager
def sqlite_reader(path):
    connection = sqlite3.connect(path.resolve().as_uri() + "?mode=ro&immutable=1", uri=True)
    try:
        connection.enable_load_extension(False)
        connection.execute("PRAGMA query_only=ON")
        connection.execute("PRAGMA trusted_schema=OFF")
        connection.setlimit(sqlite3.SQLITE_LIMIT_LENGTH, 10 * 1024 * 1024)
        deadline = time.monotonic() + 15
        connection.set_progress_handler(lambda: int(time.monotonic() > deadline), 1000)
        yield connection
    finally:
        connection.close()


def sqlite_tables(connection):
    return [row[0] for row in connection.execute(
        "SELECT name FROM sqlite_schema WHERE type='table' AND name NOT LIKE 'sqlite_%' "
        "AND upper(ltrim(sql)) NOT LIKE 'CREATE VIRTUAL%' ORDER BY name")]


def inspect_file(path: Path, filename: str, cap: int, member: str | None = None):
    try:
        if Path(filename).suffix.lower() == ".zip" and not member:
            with checked_zip(path, cap) as archive:
                members = [i.filename for i in archive.infolist() if not i.is_dir()
                           and Path(i.filename).suffix.lower() in (TEXT_EXT | STRUCTURED_EXT) - {".zip"}]
            if not members:
                reject("ZIP contains no supported dataset files")
            if len(members) != len(set(members)):
                reject("Archive contains duplicate filenames")
            return {"members": members, "tables": []}
        with source_file(path, filename, cap, member) as (source, name):
            ext = Path(name).suffix.lower()
            if ext == ".xlsx":
                import openpyxl
                with checked_zip(source, cap):
                    pass
                book = openpyxl.load_workbook(source, read_only=True, data_only=False, keep_links=False)
                try:
                    tables = book.sheetnames
                finally:
                    book.close()
            elif ext in {".sqlite", ".sqlite3", ".db"}:
                with sqlite_reader(source) as connection:
                    tables = sqlite_tables(connection)
            elif ext in TEXT_EXT | {".parquet"}:
                tables = []
            else:
                reject("Unsupported structured format")
            return {"members": [], "tables": tables}
    except QuerySentinelError:
        raise
    except Exception:
        reject("Cannot read this file; check its format and whether it is encrypted or damaged")


def choose(choices, selection):
    if selection is None and len(choices) == 1:
        return choices[0]
    if selection not in choices:
        reject("Select a worksheet or table before previewing this file")
    return selection


def cell(value):
    if value is None:
        return ""
    if isinstance(value, bytes):
        reject("Binary/blob columns are not supported; export readable values instead")
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, default=cell)
    if isinstance(value, float) and not math.isfinite(value):
        return ""
    if isinstance(value, (str, int, float, bool, decimal.Decimal)):
        return str(value)
    reject("Unsupported cell type")


def convert_file(path: Path, filename: str, destination: Path, cap: int,
                 selection: str | None = None, member: str | None = None):
    """Return (canonical filename, notes). No uploaded SQL, formulas or macros execute."""
    try:
        with source_file(path, filename, cap, member) as (source, name):
            ext = Path(name).suffix.lower()
            if ext in TEXT_EXT:
                destination.write_bytes(source.read_bytes())
                return "selected" + ext, [f"Selected archive file: {name}"]
            with contextlib.ExitStack() as stack:
                if ext == ".xlsx":
                    import openpyxl
                    with checked_zip(source, cap):
                        pass
                    book = openpyxl.load_workbook(source, read_only=True, data_only=False, keep_links=False)
                    stack.callback(book.close)
                    sheet = book[choose(book.sheetnames, selection)]
                    if sheet.max_column and sheet.max_column > MAX_COLUMNS:
                        reject("Worksheet exceeds 512 columns")
                    def excel_rows():
                        for row in sheet.iter_rows():
                            if any(c.data_type in {"f", "e"} for c in row):
                                reject("Worksheet contains formulas or error cells; export values only first")
                            yield [c.value for c in row]
                    rows = iter(excel_rows())
                    headers = next(rows, [])
                elif ext == ".parquet":
                    import pyarrow.parquet as pq
                    parquet = pq.ParquetFile(source)
                    stack.callback(parquet.close)
                    if parquet.metadata.num_rows > MAX_ROWS:
                        reject("Dataset exceeds one million rows")
                    if sum(parquet.metadata.row_group(i).total_byte_size
                           for i in range(parquet.num_row_groups)) > cap:
                        reject("Parquet expands beyond the upload limit")
                    headers = parquet.schema_arrow.names
                    rows = ([batch.column(i)[j].as_py() for i in range(batch.num_columns)]
                            for batch in parquet.iter_batches(batch_size=1024)
                            for j in range(batch.num_rows))
                elif ext in {".sqlite", ".sqlite3", ".db"}:
                    connection = stack.enter_context(sqlite_reader(source))
                    table = choose(sqlite_tables(connection), selection)
                    # Table name is selected from the file's catalog and quoted as an identifier.
                    connection.set_authorizer(lambda action, *args: sqlite3.SQLITE_OK
                                              if action in {sqlite3.SQLITE_SELECT, sqlite3.SQLITE_READ}
                                              else sqlite3.SQLITE_DENY)
                    cursor = connection.execute('SELECT * FROM "' + table.replace('"', '""') + '"')
                    headers = [c[0] for c in cursor.description]
                    rows = cursor
                else:
                    reject("Unsupported structured format")
                if not headers or len(headers) > MAX_COLUMNS:
                    reject("Dataset must have between 1 and 512 columns")
                with destination.open("w", encoding="utf-8", newline="") as output:
                    writer = csv.writer(output)
                    writer.writerow([cell(h) if h is not None else f"column_{i+1}" for i, h in enumerate(headers)])
                    for index, row in enumerate(rows, 1):
                        if index > MAX_ROWS:
                            reject("Dataset exceeds one million rows")
                        writer.writerow([cell(v) for v in row])
                        if output.tell() > cap:
                            reject("Converted dataset exceeds the upload limit")
                return "converted.csv", [f"Imported {name}" + (f" / {selection}" if selection else ""),
                    "Types are inferred from values. Workbook formatting and database indexes/constraints are not imported."]
    except QuerySentinelError:
        destination.unlink(missing_ok=True)
        raise
    except Exception:
        destination.unlink(missing_ok=True)
        reject("Cannot convert this file; check its format, cell types and whether it is encrypted or damaged")
